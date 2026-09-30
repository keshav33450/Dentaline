#!/usr/bin/env python3
"""
Phase VII - statistics exactly as in the proforma, computed on the OUT-OF-FOLD predictions
(each tooth scored by a model that never saw that patient).

Severity (No Caries / Mild / Moderate / Advanced) vs ICDAS reference:
  per class (one-vs-rest): sensitivity, specificity, PPV, NPV, accuracy, F1, ROC-AUC
  overall: accuracy, Cohen's kappa + linear-weighted kappa (ordinal)  -> 95% CI by PATIENT-level bootstrap
  confusion matrix, ROC curves, binary screening view (any caries vs none)
Tooth type (premolar vs molar): accuracy, sensitivity, specificity, PPV, NPV, kappa, AUC (+95% CI)
Model comparison: EfficientNet-B0 vs MobileNetV3 (McNemar test on the same teeth)

Outputs -> <project>/results/occlusal/
  OCCLUSAL_REPORT.md, occlusal_statistics.xlsx (per-tooth sheet ready for SPSS + all tables),
  confusion_matrix_<arch>.png, roc_<arch>.png, metrics.json
"""
from __future__ import annotations

import argparse
import csv
import json
import sys
from pathlib import Path

import numpy as np

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))
from dentassist import paths as P  # noqa: E402
from dentassist.occlusal.labels import SEVERITY, SEVERITY_LABEL, TOOTH_TYPES  # noqa: E402

RES = P.RESULTS / "occlusal"


def load(stage, arch):
    f = RES / stage / f"oof_{arch}.csv"
    if not f.exists():
        return None
    with open(f, encoding="utf-8") as fh:
        return list(csv.DictReader(fh))


def ovr(y, p, c):
    tp = int(np.sum((p == c) & (y == c))); fp = int(np.sum((p == c) & (y != c)))
    fn = int(np.sum((p != c) & (y == c))); tn = int(np.sum((p != c) & (y != c)))
    d = lambda a, b: a / b if b else float("nan")
    sens, spec, ppv, npv = d(tp, tp + fn), d(tn, tn + fp), d(tp, tp + fp), d(tn, tn + fn)
    return {"TP": tp, "FP": fp, "FN": fn, "TN": tn, "sensitivity": sens, "specificity": spec, "PPV": ppv, "NPV": npv,
            "accuracy": d(tp + tn, tp + tn + fp + fn), "F1": d(2 * tp, 2 * tp + fp + fn)}


def kappa(y, p, k, weights=None):
    from sklearn.metrics import cohen_kappa_score
    if len(set(y) | set(p)) < 2:
        return float("nan")
    return float(cohen_kappa_score(y, p, labels=list(range(k)), weights=weights))


def auc_ovr(y, probs, c):
    from sklearn.metrics import roc_auc_score
    t = (y == c).astype(int)
    return float(roc_auc_score(t, probs[:, c])) if 0 < t.sum() < len(t) else float("nan")


def boot_ci(fn, y, p, probs, pats, n=1000, seed=0):
    """patient-level (cluster) bootstrap 95% CI"""
    rng = np.random.default_rng(seed)
    up = np.unique(pats)
    idx = {u: np.where(pats == u)[0] for u in up}
    vals = []
    for _ in range(n):
        s = np.concatenate([idx[u] for u in rng.choice(up, len(up), replace=True)])
        v = fn(y[s], p[s], probs[s])
        if not np.isnan(v):
            vals.append(v)
    return (float(np.percentile(vals, 2.5)), float(np.percentile(vals, 97.5))) if len(vals) > 20 else (float("nan"),) * 2


def analyse(rows, classes, n_boot):
    k = len(classes)
    y = np.array([classes.index(r["true"]) for r in rows]); p = np.array([classes.index(r["pred"]) for r in rows])
    probs = np.array([[float(r[f"p_{c}"]) for c in classes] for r in rows]); pats = np.array([r["patient_id"] for r in rows])
    cm = np.zeros((k, k), int)
    for a, b in zip(y, p):
        cm[a, b] += 1
    per = {}
    for c in range(k):
        m = ovr(y, p, c)
        m["AUC"] = auc_ovr(y, probs, c)
        for key, f in (("sensitivity", lambda yy, pp, pr, c=c: ovr(yy, pp, c)["sensitivity"]),
                       ("specificity", lambda yy, pp, pr, c=c: ovr(yy, pp, c)["specificity"]),
                       ("AUC", lambda yy, pp, pr, c=c: auc_ovr(yy, pr, c))):
            m[key + "_95CI"] = boot_ci(f, y, p, probs, pats, n_boot)
        per[classes[c]] = m
    overall = {
        "n_teeth": int(len(y)), "n_patients": int(len(np.unique(pats))),
        "accuracy": float((y == p).mean()),
        "accuracy_95CI": boot_ci(lambda a, b, _: float((a == b).mean()), y, p, probs, pats, n_boot),
        "kappa": kappa(y, p, k), "kappa_95CI": boot_ci(lambda a, b, _: kappa(a, b, k), y, p, probs, pats, n_boot),
        "macro_F1": float(np.nanmean([per[c]["F1"] for c in classes])),
    }
    if k > 2:
        overall["weighted_kappa_linear"] = kappa(y, p, k, "linear")
        overall["weighted_kappa_linear_95CI"] = boot_ci(lambda a, b, _: kappa(a, b, k, "linear"), y, p, probs, pats, n_boot)
        yb, pb = (y > 0).astype(int), (p > 0).astype(int)
        scr = ovr(yb, pb, 1)
        scr["AUC"] = auc_ovr(yb, np.stack([probs[:, 0], 1 - probs[:, 0]], 1), 1)
        overall["screening_any_caries_vs_none"] = scr
    return {"overall": overall, "per_class": per, "confusion": cm.tolist(), "y": y, "p": p, "probs": probs}


def mcnemar(rows_a, rows_b):
    ca = {r["crop"]: r["true"] == r["pred"] for r in rows_a}; cb = {r["crop"]: r["true"] == r["pred"] for r in rows_b}
    common = ca.keys() & cb.keys()
    b = sum(ca[c] and not cb[c] for c in common); c = sum(cb[c] and not ca[c] for c in common)
    from scipy.stats import binomtest
    pval = binomtest(min(b, c), b + c, 0.5).pvalue if b + c else 1.0
    return {"only_A_correct": b, "only_B_correct": c, "p_value_exact": float(pval)}


def plots(res, classes, tag):
    import matplotlib
    matplotlib.use("Agg")
    import matplotlib.pyplot as plt
    from sklearn.metrics import roc_curve
    names = [SEVERITY_LABEL.get(c, c.title()) for c in classes]
    cm = np.array(res["confusion"])
    fig, ax = plt.subplots(figsize=(5.2, 4.4))
    ax.imshow(cm, cmap="Blues")
    for i in range(len(classes)):
        for j in range(len(classes)):
            ax.text(j, i, cm[i, j], ha="center", va="center", color="white" if cm[i, j] > cm.max() / 2 else "black")
    ax.set_xticks(range(len(classes)), names, rotation=30); ax.set_yticks(range(len(classes)), names)
    ax.set_xlabel("AI prediction"); ax.set_ylabel("ICDAS reference"); ax.set_title(f"Confusion matrix - {tag}")
    fig.tight_layout(); fig.savefig(RES / f"confusion_matrix_{tag}.png", dpi=200); plt.close(fig)
    fig, ax = plt.subplots(figsize=(5, 4.6))
    for c, nm in enumerate(names):
        t = (res["y"] == c).astype(int)
        if 0 < t.sum() < len(t):
            fpr, tpr, _ = roc_curve(t, res["probs"][:, c])
            ax.plot(fpr, tpr, label=f"{nm} (AUC {res['per_class'][classes[c]]['AUC']:.2f})")
    ax.plot([0, 1], [0, 1], "--", color="grey"); ax.set_xlabel("1 - specificity"); ax.set_ylabel("Sensitivity")
    ax.set_title(f"ROC - {tag}"); ax.legend(fontsize=8); fig.tight_layout(); fig.savefig(RES / f"roc_{tag}.png", dpi=200); plt.close(fig)


def fmt(v, pct=True):
    if v is None or (isinstance(v, float) and np.isnan(v)):
        return "-"
    return f"{v * 100:.1f}%" if pct else f"{v:.3f}"


def ci(t, pct=True):
    return f"({fmt(t[0], pct)}-{fmt(t[1], pct)})" if t and not np.isnan(t[0]) else ""


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--arch", nargs="+", default=["efficientnet_b0", "mobilenet_v3"])
    ap.add_argument("--bootstrap", type=int, default=1000)
    a = ap.parse_args()
    RES.mkdir(parents=True, exist_ok=True)
    out, md, sheets = {}, ["# Occlusal caries AI - diagnostic performance (out-of-fold, patient-grouped CV)", ""], {}
    selected = {}
    for stage, classes in (("severity", SEVERITY), ("type", TOOTH_TYPES)):
        sel_f = RES / stage / "cv_summary.json"
        if sel_f.exists():
            selected[stage] = json.loads(sel_f.read_text()).get("selected_for_app")
        loaded = {ar: load(stage, ar) for ar in a.arch}
        loaded = {k: v for k, v in loaded.items() if v}
        for arch, rows in loaded.items():
            res = analyse(rows, classes, a.bootstrap)
            tag = f"{stage}_{arch}"
            plots(res, classes, tag)
            out[tag] = {k: v for k, v in res.items() if k not in ("y", "p", "probs")}
            o = res["overall"]
            md += [f"## {'Caries severity' if stage == 'severity' else 'Tooth type'} - {arch}"
                   + ("  (selected for app)" if selected.get(stage) == arch else ""), "",
                   f"Teeth {o['n_teeth']}, patients {o['n_patients']}. Accuracy **{fmt(o['accuracy'])}** {ci(o['accuracy_95CI'])}, "
                   f"Cohen's kappa **{fmt(o['kappa'], False)}** {ci(o['kappa_95CI'], False)}"
                   + (f", linear-weighted kappa **{fmt(o['weighted_kappa_linear'], False)}** {ci(o['weighted_kappa_linear_95CI'], False)}" if "weighted_kappa_linear" in o else "")
                   + f", macro-F1 {fmt(o['macro_F1'], False)}", "",
                   "| Class | Sensitivity (95% CI) | Specificity (95% CI) | PPV | NPV | Accuracy | F1 | AUC (95% CI) |",
                   "|---|---|---|---|---|---|---|---|"]
            for c in classes:
                m = res["per_class"][c]
                md.append(f"| {SEVERITY_LABEL.get(c, c.title())} | {fmt(m['sensitivity'])} {ci(m['sensitivity_95CI'])} | "
                          f"{fmt(m['specificity'])} {ci(m['specificity_95CI'])} | {fmt(m['PPV'])} | {fmt(m['NPV'])} | "
                          f"{fmt(m['accuracy'])} | {fmt(m['F1'], False)} | {fmt(m['AUC'], False)} {ci(m['AUC_95CI'], False)} |")
            if "screening_any_caries_vs_none" in o:
                s = o["screening_any_caries_vs_none"]
                md += ["", f"Screening view (any caries vs sound): sensitivity {fmt(s['sensitivity'])}, specificity {fmt(s['specificity'])}, "
                           f"PPV {fmt(s['PPV'])}, NPV {fmt(s['NPV'])}, AUC {fmt(s['AUC'], False)}"]
            md += ["", f"![cm](confusion_matrix_{tag}.png) ![roc](roc_{tag}.png)", ""]
            sheets[f"{stage}_{arch}"[:31]] = (["Class", "TP", "FP", "FN", "TN", "Sensitivity", "Sens 95%CI", "Specificity",
                                               "Spec 95%CI", "PPV", "NPV", "Accuracy", "F1", "AUC", "AUC 95%CI"],
                                              [[SEVERITY_LABEL.get(c, c), m["TP"], m["FP"], m["FN"], m["TN"], m["sensitivity"],
                                                ci(m["sensitivity_95CI"]), m["specificity"], ci(m["specificity_95CI"]), m["PPV"], m["NPV"],
                                                m["accuracy"], m["F1"], m["AUC"], ci(m["AUC_95CI"], False)]
                                               for c, m in res["per_class"].items()]
                                              + [[], ["Overall accuracy", o["accuracy"], ci(o["accuracy_95CI"])],
                                                 ["Cohen kappa", o["kappa"], ci(o["kappa_95CI"], False)]]
                                              + ([["Linear weighted kappa", o["weighted_kappa_linear"], ci(o["weighted_kappa_linear_95CI"], False)]]
                                                 if "weighted_kappa_linear" in o else []))
        if len(loaded) == 2:
            (a1, r1), (a2, r2) = loaded.items()
            mc = mcnemar(r1, r2)
            out[f"{stage}_mcnemar"] = mc
            md += [f"**{stage}: {a1} vs {a2}** - McNemar exact p = {mc['p_value_exact']:.3f} "
                   f"(only {a1} correct: {mc['only_A_correct']}, only {a2} correct: {mc['only_B_correct']})", ""]

    # per-tooth sheet (SPSS-ready) + proforma match/mismatch table, from the app-selected models
    sev = {r["crop"]: r for r in (load("severity", selected.get("severity", a.arch[0])) or [])}
    typ = {r["crop"]: r for r in (load("type", selected.get("type", a.arch[0])) or [])}
    head = ["patient_id", "image", "fdi", "icdas_code", "ref_severity", "ai_severity", "ai_severity_conf", "severity_match",
            "ref_type", "ai_type", "ai_type_conf", "type_match"]
    per_tooth = []
    for crop in sorted(sev.keys() | typ.keys()):
        s, t = sev.get(crop, {}), typ.get(crop, {})
        base = s or t
        per_tooth.append([base.get("patient_id"), base.get("image"), base.get("fdi"), s.get("icdas", ""),
                          SEVERITY_LABEL.get(s.get("true"), ""), SEVERITY_LABEL.get(s.get("pred"), ""), s.get("confidence", ""),
                          ("Yes" if s.get("true") == s.get("pred") else "No") if s else "",
                          t.get("true", ""), t.get("pred", ""), t.get("confidence", ""),
                          ("Yes" if t.get("true") == t.get("pred") else "No") if t else ""])
    sheets = {"per_tooth (SPSS)": (head, per_tooth), **sheets}
    prof = [["Tooth Type", "Premolar / Molar",
             sum(r[9] == r[8] for r in per_tooth if r[8]), sum(r[9] != r[8] for r in per_tooth if r[8])]]
    for c in SEVERITY:
        L = SEVERITY_LABEL[c]
        rr = [r for r in per_tooth if r[4] == L]
        prof.append(["Caries Severity", L, sum(r[5] == L for r in rr), sum(r[5] != L for r in rr)])
    sheets["proforma_table"] = (["Variable Assessed", "ICDAS Reference Standard", "AI match (Yes)", "AI mismatch (No)"], prof)
    md += ["## Proforma evaluation table (AI vs ICDAS)", "", "| Variable | ICDAS reference | Match (Yes) | Mismatch (No) |", "|---|---|---|---|"]
    md += [f"| {r[0]} | {r[1]} | {r[2]} | {r[3]} |" for r in prof]

    import openpyxl
    from openpyxl.styles import Font
    wb = openpyxl.Workbook(); wb.remove(wb.active)
    for name, (h, rows) in sheets.items():
        ws = wb.create_sheet(name[:31]); ws.append(h)
        for c in ws[1]:
            c.font = Font(bold=True)
        for r in rows:
            ws.append([None if (isinstance(v, float) and np.isnan(v)) else v for v in r])
    wb.save(RES / "occlusal_statistics.xlsx")
    (RES / "metrics.json").write_text(json.dumps(out, indent=2, default=lambda x: None))
    (RES / "OCCLUSAL_REPORT.md").write_text("\n".join(md) + "\n", encoding="utf-8")
    print("\n".join(md))
    print(f"\n[done] {RES}/OCCLUSAL_REPORT.md, occlusal_statistics.xlsx")


if __name__ == "__main__":
    main()
